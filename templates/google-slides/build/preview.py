"""Render a generated PPTX to HTML by reading its real shapes — a true preview."""
import os, sys, html
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
PX = 96.0          # px per inch in the preview
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def px(v):
    return v / EMU * PX


def color_of(fmt):
    try:
        if fmt.type is not None and fmt.fore_color.type is not None:
            return "#%02X%02X%02X" % fmt.fore_color.rgb[0:3] if False else "#" + str(fmt.fore_color.rgb)
    except Exception:
        pass
    return None


def fill_css(sh):
    from pptx.enum.dml import MSO_FILL_TYPE
    try:
        f = sh.fill
        if f.type == MSO_FILL_TYPE.SOLID:
            return "#" + str(f.fore_color.rgb)
    except Exception:
        pass
    return None


def line_css(sh):
    try:
        c = "#" + str(sh.line.color.rgb)
        w = sh.line.width.pt if sh.line.width else 1
        return c, w
    except Exception:
        return None, 0


def custgeom_points(sh):
    """Extract freeform path points (local EMU) if present."""
    el = sh._element
    paths = el.findall('.//a:custGeom/a:pathLst/a:path', NS)
    out = []
    for p in paths:
        pw = int(p.get('w') or 0); ph = int(p.get('h') or 0)
        pts = []
        for child in p:
            tag = child.tag.split('}')[1]
            if tag in ('moveTo', 'lnTo'):
                pt = child.find('a:pt', NS)
                pts.append((int(pt.get('x')), int(pt.get('y'))))
            elif tag == 'close':
                pts.append(None)
        out.append((pw, ph, pts))
    return out


def render(path, out_html, max_slides=None):
    prs = Presentation(path)
    SW, SH = px(prs.slide_width), px(prs.slide_height)
    parts = ["<html><head><meta charset='utf-8'><style>",
             "body{margin:0;background:#888;font-family:Inter,sans-serif}",
             ".s{position:relative;overflow:hidden;margin:10px auto;box-shadow:0 0 0 1px #333;",
             f"width:{SW}px;height:{SH}px;background:#fff}}",
             ".n{position:absolute;left:4px;top:4px;font:10px monospace;color:#999;z-index:99}",
             "</style></head><body>"]
    for idx, sl in enumerate(prs.slides, 1):
        if max_slides and idx > max_slides:
            break
        parts.append(f"<div class='s'><div class='n'>{idx}</div>")
        for sh in sl.shapes:
            if sh.left is None:
                continue
            L, T, W, H = px(sh.left), px(sh.top), px(abs(sh.width)), px(abs(sh.height))
            st = f"position:absolute;left:{L:.1f}px;top:{T:.1f}px;width:{W:.1f}px;height:{H:.1f}px;"
            geo = custgeom_points(sh)
            if geo:
                lc, lw = line_css(sh)
                fc = fill_css(sh)
                for pw, ph, pts in geo:
                    if not pts or pw == 0 or ph == 0:
                        continue
                    d = []
                    started = False
                    for pt in pts:
                        if pt is None:
                            d.append("Z"); continue
                        x = pt[0] / pw * W
                        y = pt[1] / ph * H
                        d.append(("M" if not started else "L") + f"{x:.1f},{y:.1f}")
                        started = True
                    parts.append(
                        f"<svg style='{st}' viewBox='0 0 {W:.1f} {H:.1f}'><path d='{' '.join(d)}' "
                        f"fill='{fc or 'none'}' stroke='{lc or 'none'}' stroke-width='{lw*1.33:.1f}' "
                        f"stroke-linecap='square' stroke-linejoin='miter'/></svg>")
                continue
            if sh.shape_type is not None and 'CONNECTOR' in str(sh.shape_type):
                lc, lw = line_css(sh)
                flipH = sh._element.find('.//a:xfrm', NS)
                fh = flipH is not None and flipH.get('flipH') == '1'
                fv = flipH is not None and flipH.get('flipV') == '1'
                x1, y1, x2, y2 = (W, 0, 0, H) if fh and not fv else \
                                 (0, H, W, 0) if fv and not fh else (0, 0, W, H)
                parts.append(f"<svg style='{st}' viewBox='0 0 {max(W,1):.1f} {max(H,1):.1f}'>"
                             f"<line x1='{x1:.1f}' y1='{y1:.1f}' x2='{x2:.1f}' y2='{y2:.1f}' "
                             f"stroke='{lc or '#000'}' stroke-width='{lw*1.33:.1f}' stroke-linecap='square'/></svg>")
                continue
            bg = fill_css(sh)
            lc, lw = line_css(sh)
            box = st
            if bg:
                box += f"background:{bg};"
            if lc and lw:
                box += f"border:{lw*1.33:.1f}px solid {lc};box-sizing:border-box;"
            if 'OVAL' in str(sh.shape_type):
                box += "border-radius:50%;"
            inner = ""
            if sh.has_text_frame and sh.text_frame.text.strip():
                lines = []
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs)
                    if not t:
                        lines.append("<br>"); continue
                    r0 = para.runs[0]
                    fs = (r0.font.size.pt if r0.font.size else 13.5) * 1.3333
                    fam = r0.font.name or "Inter"
                    col = "#" + str(r0.font.color.rgb) if r0.font.color and r0.font.color.type is not None else "#000"
                    bold = "700" if r0.font.bold else "400"
                    al = {2: "center", 3: "right"}.get(
                        para.alignment.value if para.alignment else 0, "left")
                    ls = para.line_spacing or 1.15
                    lines.append(
                        f"<div style=\"font-family:'{fam}',sans-serif;font-size:{fs}px;color:{col};"
                        f"font-weight:{bold};text-align:{al};line-height:{ls};white-space:pre-wrap\">"
                        f"{html.escape(t)}</div>")
                inner = "".join(lines)
            parts.append(f"<div style='{box}'>{inner}</div>")
        parts.append("</div>")
    parts.append("</body></html>")
    open(out_html, "w").write("\n".join(parts))
    print("wrote", out_html, "slides:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else None)
