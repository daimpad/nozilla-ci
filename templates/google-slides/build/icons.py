"""nozilla icons — one source, two outputs.

Every icon is defined once on a 64-unit grid with a 4-unit stroke (square caps,
miter joins) and can be rendered either as native PPTX vector shapes (sharp and
recolourable inside Google Slides) or exported as SVG for the web — so the deck
and project/assets/ can never drift apart.
"""
try:
    from pptx.util import Pt, Emu
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
except ImportError:  # SVG-Export und Pruefung kommen ohne python-pptx aus
    class RGBColor(tuple):
        def __new__(cls, r, g, b):
            return super().__new__(cls, (r, g, b))

GREEN = RGBColor(0x00, 0xFF, 0x9C)
GREEN_DEEP = RGBColor(0x00, 0xC0, 0x75)
GREEN_SOFT = RGBColor(0xB7, 0xFF, 0xE0)
G = 64.0


def _emu(v):
    return Emu(int(round(v)))


# ── Renderer A: native PPTX shapes ────────────────────────────────────────
class Pen:
    def __init__(self, slide, x, y, size, color, stroke=4.0):
        self.s = slide
        self.x0, self.y0 = x * 914400.0, y * 914400.0
        self.k = size * 914400.0 / G
        self.color = color
        self.wpt = Pt(size * (stroke / G) * 72.0)

    def X(self, u): return self.x0 + u * self.k
    def Y(self, v): return self.y0 + v * self.k

    def _plain(self, sh):
        sh.shadow.inherit = False
        return sh

    def line(self, x1, y1, x2, y2):
        from pptx.enum.shapes import MSO_CONNECTOR
        c = self.s.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, _emu(self.X(x1)), _emu(self.Y(y1)),
            _emu(self.X(x2)), _emu(self.Y(y2)))
        c.line.color.rgb = self.color
        c.line.width = self.wpt
        return self._plain(c)

    def rect(self, x, y, w, h, fill=None):
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(self.X(x)), _emu(self.Y(y)),
                                     _emu(w * self.k), _emu(h * self.k))
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = self.color
        sh.line.width = self.wpt
        return self._plain(sh)

    def solid(self, x, y, w, h, fill):
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(self.X(x)), _emu(self.Y(y)),
                                     _emu(w * self.k), _emu(h * self.k))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        return self._plain(sh)

    def circle(self, cx, cy, r, fill=None):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, _emu(self.X(cx - r)), _emu(self.Y(cy - r)),
                                     _emu(2 * r * self.k), _emu(2 * r * self.k))
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = self.color
        sh.line.width = self.wpt
        return self._plain(sh)

    def dot(self, cx, cy, r, fill):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, _emu(self.X(cx - r)), _emu(self.Y(cy - r)),
                                     _emu(2 * r * self.k), _emu(2 * r * self.k))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        return self._plain(sh)

    def poly(self, pts, close=False, fill=None):
        ff = self.s.shapes.build_freeform(_emu(self.X(pts[0][0])), _emu(self.Y(pts[0][1])))
        ff.add_line_segments([(_emu(self.X(x)), _emu(self.Y(y))) for x, y in pts[1:]], close=close)
        sh = ff.convert_to_shape()
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = self.color
        sh.line.width = self.wpt
        return self._plain(sh)


# ── Renderer B: SVG ───────────────────────────────────────────────────────
def _hex(c):
    return "#%02X%02X%02X" % (c[0], c[1], c[2]) if not isinstance(c, str) else c


class SvgPen:
    """Same drawing API, emits SVG — keeps web assets identical to the deck."""

    def __init__(self, color="currentColor", stroke=4.0):
        self.color = color
        self.sw = stroke
        self.out = []

    def line(self, x1, y1, x2, y2):
        self.out.append(f'<path d="M{x1} {y1} L{x2} {y2}"/>')

    def rect(self, x, y, w, h, fill=None):
        f = f' fill="{_hex(fill)}"' if fill is not None else ''
        self.out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}"{f}/>')

    def solid(self, x, y, w, h, fill):
        self.out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" '
                        f'fill="{_hex(fill)}" stroke="none"/>')

    def circle(self, cx, cy, r, fill=None):
        f = f' fill="{_hex(fill)}"' if fill is not None else ''
        self.out.append(f'<circle cx="{cx}" cy="{cy}" r="{r}"{f}/>')

    def dot(self, cx, cy, r, fill):
        self.out.append(f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="{_hex(fill)}" stroke="none"/>')

    def poly(self, pts, close=False, fill=None):
        d = "M" + " L".join(f"{x} {y}" for x, y in pts) + (" Z" if close else "")
        f = f' fill="{_hex(fill)}"' if fill is not None else ''
        self.out.append(f'<path d="{d}"{f}/>')

    def svg(self, label):
        body = "\n  ".join(self.out)
        return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64" fill="none" '
                f'stroke="{self.color}" stroke-width="{self.sw}" stroke-linecap="square" '
                f'stroke-linejoin="miter" role="img" aria-label="{label}">\n  {body}\n</svg>\n')


# ── Dialekt A — Heavy-Stroke ──────────────────────────────────────────────
def i_grid(p):
    p.rect(9, 9, 17, 17, fill=GREEN); p.rect(30, 9, 17, 17)
    p.rect(9, 30, 17, 17); p.rect(30, 30, 17, 17)


def i_cursor(p):
    """Ein Fenster, darin ein Zeiger. Die Webseite zeigt damit Bedienung, nicht
    nur den Mauszeiger."""
    p.rect(7, 12, 45, 30); p.line(7, 48, 52, 48)
    p.solid(14, 19, 15, 9, GREEN)
    p.poly([(34, 18), (34, 38), (39, 33), (43, 40), (46, 39), (42, 32), (48, 32)],
           close=True, fill=p.color)


def i_maze(p):
    p.solid(28, 22, 10, 10, GREEN)
    p.poly([(12, 12), (49, 12), (49, 49), (20, 49), (20, 20), (41, 20), (41, 41),
            (27, 41), (27, 27), (34, 27)])


def i_terminal(p):
    p.solid(8, 12, 44, 11, GREEN)
    p.rect(8, 12, 44, 38); p.line(8, 23, 52, 23)
    p.poly([(16, 31), (22, 36), (16, 41)]); p.line(28, 41, 42, 41)


def i_chat(p):
    p.poly([(10, 12), (50, 12), (50, 37), (25, 37), (16, 46), (16, 37), (10, 37)],
           close=True, fill=GREEN)
    p.dot(24, 24, 2.4, p.color); p.dot(31, 24, 2.4, p.color); p.dot(38, 24, 2.4, p.color)


def i_spark(p):
    p.poly([(31, 8), (36, 26), (54, 31), (36, 36), (31, 50), (26, 36), (8, 31), (26, 26)],
           close=True, fill=GREEN)


def i_arrow(p):
    p.line(10, 32, 50, 32); p.poly([(38, 20), (52, 32), (38, 44)])


def i_user(p):
    """Kopf und Schultern. Die Webseite zeichnet die Schultern als Bogen, hier
    sind es Kanten: derselbe Umriss, in der Sprache dieses Hauses."""
    p.circle(31, 21, 10, fill=GREEN)
    p.poly([(13, 50), (17, 40), (31, 36), (45, 40), (49, 50)])


def i_gear(p):
    p.circle(31, 31, 14); p.dot(31, 31, 5, GREEN)
    p.line(45, 31, 51, 31); p.line(31, 45, 31, 51)
    p.line(17, 31, 11, 31); p.line(31, 17, 31, 11)
    p.line(41, 41, 45, 45); p.line(21, 41, 17, 45)
    p.line(21, 21, 17, 17); p.line(41, 21, 45, 17)


def i_folder(p):
    p.poly([(10, 20), (26, 20), (31, 26), (51, 26), (51, 48), (10, 48)],
           close=True, fill=GREEN)


def i_lightbulb(p):
    p.circle(31, 25, 14, fill=GREEN)
    p.line(23, 43, 39, 43); p.line(25, 49, 37, 49)


def i_code(p):
    p.solid(26, 26, 12, 12, GREEN)
    p.poly([(24, 18), (12, 32), (24, 46)]); p.poly([(40, 18), (51, 32), (40, 46)])
    p.line(36, 14, 28, 50)


def i_calendar(p):
    p.rect(8, 14, 44, 34); p.line(8, 26, 52, 26)
    p.line(19, 8, 19, 18); p.line(41, 8, 41, 18)
    p.solid(16, 33, 10, 9, GREEN)


def i_clock(p):
    p.circle(32, 32, 21); p.line(32, 32, 32, 19); p.line(32, 32, 43, 32)


def i_chart(p):
    p.line(10, 10, 10, 50); p.line(10, 50, 50, 50)
    p.solid(16, 32, 8, 18, p.color); p.solid(28, 24, 8, 26, p.color)
    p.solid(40, 14, 8, 36, GREEN)


def i_target(p):
    p.circle(32, 32, 21); p.circle(32, 32, 11); p.dot(32, 32, 4.5, GREEN)


def i_lock(p):
    p.rect(13, 29, 38, 25)
    p.poly([(22, 29), (22, 20), (32, 13), (42, 20), (42, 29)])
    p.dot(32, 41, 3.5, GREEN)


def i_server(p):
    p.rect(10, 12, 42, 16); p.rect(10, 34, 42, 16)
    p.line(39, 20, 45, 20); p.line(39, 42, 45, 42)
    p.solid(17, 16, 8, 8, GREEN)


def i_mobile(p):
    p.rect(20, 8, 24, 48); p.line(27, 15, 37, 15); p.line(26, 49, 38, 49)


def i_monitor(p):
    p.rect(8, 12, 48, 31); p.line(32, 43, 32, 52); p.line(22, 52, 42, 52)


def i_search(p):
    p.circle(27, 27, 14); p.line(37, 37, 50, 50)


def i_mail(p):
    p.rect(8, 16, 44, 30)
    p.poly([(8, 16), (30, 34), (52, 16)], close=True, fill=GREEN)


def i_phone(p):
    """Hoerer: zwei Muscheln, dazwischen ein Buegel. Kanten statt Bogen."""
    p.poly([(10, 38), (10, 24), (22, 13), (42, 13), (54, 24), (54, 38)])
    p.rect(6, 36, 15, 13); p.solid(43, 38, 11, 9, GREEN)
    p.rect(43, 36, 15, 13)


def i_callback(p):
    """Rueckruf: derselbe Hoerer, darueber ein Pfeil, der umkehrt.

    Der Pfeil sitzt hoeher als der Hoerer und beruehrt ihn nicht: Zwei Formen,
    die sich kreuzen, liest man als eine dritte.
    """
    p.poly([(12, 42), (12, 32), (23, 25), (41, 25), (52, 32), (52, 42)])
    p.rect(7, 40, 15, 10); p.rect(42, 40, 15, 10)
    p.poly([(46, 12), (46, 5), (28, 5)])
    p.poly([(19, 6), (37, 6), (28, 19)], close=True, fill=GREEN)


def i_form(p):
    """Formular: ein Blatt mit zwei Feldern, das zweite ausgefuellt."""
    p.rect(10, 8, 38, 46)
    p.rect(18, 19, 22, 9)
    p.solid(18, 35, 22, 9, GREEN)


def i_warning(p):
    p.poly([(32, 8), (54, 50), (10, 50)], close=True)
    p.line(32, 24, 32, 37); p.dot(32, 44, 2.6, p.color)


def i_check(p):
    p.poly([(12, 34), (26, 48), (52, 16)])


def i_close(p):
    p.line(14, 14, 50, 50); p.line(50, 14, 14, 50)


def i_plus(p):
    p.line(32, 12, 32, 52); p.line(12, 32, 52, 32)


def i_external(p):
    p.poly([(10, 22), (10, 50), (40, 50), (40, 32)])
    p.line(28, 32, 45, 15)
    p.poly([(34, 10), (50, 10), (50, 26)], close=True, fill=GREEN)


def i_rocket(p):
    p.poly([(32, 8), (42, 27), (42, 45), (22, 45), (22, 27)], close=True)
    p.poly([(22, 34), (13, 51), (22, 45)]); p.poly([(42, 34), (51, 51), (42, 45)])
    p.dot(32, 25, 4.5, GREEN)


def i_pen(p):
    p.poly([(13, 51), (17, 39), (43, 13), (51, 21), (25, 47)], close=True)
    p.line(39, 17, 47, 25)


def i_trash(p):
    p.line(9, 19, 55, 19); p.rect(16, 19, 32, 35); p.rect(26, 10, 12, 9)
    p.line(27, 28, 27, 45); p.line(37, 28, 37, 45)


def i_filter(p):
    p.poly([(9, 13), (55, 13), (38, 32), (38, 53), (26, 45), (26, 32)], close=True)


def i_book(p):
    """Aufgeschlagenes Buch, eine Seite gruen. Statt der Bogenkanten der
    Webseite zwei Trapeze: Der Stift kennt keinen Bogen, und das CI sagt zur
    Formsprache Kanten."""
    p.poly([(31, 16), (10, 12), (10, 45), (31, 49)], close=True)
    p.poly([(31, 16), (50, 12), (50, 45), (31, 49)], close=True, fill=GREEN)
    p.line(31, 16, 31, 49)


def i_bell(p):
    p.poly([(20, 43), (20, 28), (32, 16), (44, 28), (44, 43)], close=True)
    p.line(13, 43, 51, 43); p.dot(32, 50, 3.4, GREEN)


# ── Dialekt B — Pixel-Art ─────────────────────────────────────────────────
_CUKE_BODY = [(40, 12), (44, 12), (36, 16), (40, 16), (44, 16), (32, 20), (36, 20), (40, 20),
              (28, 24), (32, 24), (36, 24), (24, 28), (28, 28), (32, 28), (20, 32), (24, 32),
              (28, 32), (16, 36), (20, 36), (24, 36), (12, 40), (16, 40), (20, 40), (12, 44),
              (16, 44)]
_CUKE_SEED = [(40, 16), (32, 24), (24, 32), (16, 40)]


def i_cucumber(p):
    for x, y in _CUKE_BODY:
        p.solid(x, y, 4, 4, GREEN)
    for x, y in _CUKE_SEED:
        p.solid(x, y, 4, 4, GREEN_DEEP)
    for x, y in ((44, 12), (36, 20)):
        p.solid(x, y, 4, 4, GREEN_SOFT)
    p.solid(12, 44, 4, 4, p.color)


_HEART_ROWS = [(8, [1, 2, 6, 7]), (14, [0, 1, 2, 3, 5, 6, 7, 8]),
               (20, [0, 1, 2, 3, 4, 5, 6, 7, 8]), (26, [0, 1, 2, 3, 4, 5, 6, 7, 8]),
               (32, [1, 2, 3, 4, 5, 6, 7]), (38, [2, 3, 4, 5, 6]), (44, [3, 4, 5]), (50, [4])]


def i_pixel_heart(p):
    for y, cols in _HEART_ROWS:
        for c in cols:
            p.solid(5 + c * 6, y, 6, 6, GREEN)


_STAR_ROWS = [(8, [4]), (14, [3, 4, 5]), (20, [0, 1, 2, 3, 4, 5, 6, 7, 8]),
              (26, [1, 2, 3, 4, 5, 6, 7]), (32, [2, 3, 4, 5, 6]), (38, [1, 2, 3, 5, 6, 7]),
              (44, [0, 1, 7, 8])]


def i_pixel_star(p):
    for y, cols in _STAR_ROWS:
        for c in cols:
            p.solid(5 + c * 6, y, 6, 6, GREEN)


def i_pixel_disk(p):
    """Pixel-Diskette — retro save."""
    for y in range(10, 52, 6):
        for c in range(0, 8):
            p.solid(8 + c * 6, y, 6, 6, GREEN)
    for c in range(2, 6):                      # label window
        p.solid(8 + c * 6, 34, 6, 6, GREEN_SOFT)
        p.solid(8 + c * 6, 40, 6, 6, GREEN_SOFT)
    for c in range(2, 6):                      # shutter
        p.solid(8 + c * 6, 10, 6, 6, GREEN_DEEP)
        p.solid(8 + c * 6, 16, 6, 6, GREEN_DEEP)


# ── Registry: slug · label · draw ─────────────────────────────────────────
# ── Dialekt C — Konzept, zweite Staffel ───────────────────────────────────
# Betrieb
def i_ops_pulse(p):
    """Monitoring — der Ausschlag, auf den man wartet."""
    p.poly([(4, 32), (14, 32), (19, 18), (25, 46), (31, 24), (36, 32), (46, 32)])
    p.dot(48, 32, 4, GREEN)


def i_ops_incident(p):
    """Stoerung — der Blitz im Kasten."""
    p.rect(7, 11, 40, 34)
    p.poly([(30, 15), (19, 31), (27, 31), (22, 43), (35, 26), (27, 26)],
           close=True, fill=GREEN)


def i_ops_rollback(p):
    """Zuruecknehmen — die Runde rueckwaerts."""
    p.poly([(22, 13), (48, 13), (48, 46), (14, 46), (14, 27)])
    p.poly([(7, 28), (14, 17), (21, 28)], close=True, fill=GREEN)


def i_ops_scale(p):
    """Skalieren — Stufe fuer Stufe, nicht auf einmal."""
    p.line(5, 50, 49, 50)
    p.solid(8, 38, 10, 12, p.color)
    p.solid(22, 28, 10, 22, p.color)
    p.solid(36, 14, 10, 36, GREEN)


def i_ops_backup(p):
    """Sicherung — dieselbe Sache, an einem zweiten Ort."""
    p.rect(5, 13, 20, 14)
    p.rect(5, 33, 20, 14)
    p.rect(31, 23, 18, 14, fill=GREEN)
    p.line(25, 30, 31, 30)


# Sicherheit
def i_sec_encrypt(p):
    """Verschluesseln — was links steht, ist rechts nicht mehr lesbar."""
    p.rect(5, 15, 19, 30)
    for x, y, g in ((31, 16, 0), (40, 16, 1), (34, 25, 0), (44, 27, 0),
                    (31, 34, 1), (42, 36, 0), (35, 43, 0)):
        p.solid(x, y, 6, 6, GREEN if g else p.color)


def i_sec_surface(p):
    """Angriffsflaeche — alles, was von aussen erreichbar ist."""
    p.rect(17, 17, 31, 28)
    for y in (21, 31, 41):
        p.poly([(4, y - 4), (13, y), (4, y + 4)], close=True, fill=GREEN)


def i_sec_key(p):
    """Berechtigung — ein Schluessel, kein Generalschluessel."""
    p.circle(16, 31, 9)
    p.dot(16, 31, 3.4, GREEN)
    p.line(25, 31, 50, 31)
    p.line(42, 31, 42, 40)
    p.line(48, 31, 48, 37)


# Barrierefreiheit
def i_a11y_contrast(p):
    """Kontrast — der Unterschied muss messbar sein."""
    p.rect(8, 17, 22, 28, fill=p.color)
    p.rect(30, 17, 22, 28, fill=GREEN)


def i_a11y_keyboard(p):
    """Tastaturpfad — jede Station muss erreichbar sein."""
    p.rect(6, 25, 12, 12)
    p.rect(24, 25, 12, 12, fill=GREEN)
    p.rect(42, 25, 12, 12)
    p.line(18, 31, 24, 31)
    p.line(36, 31, 42, 31)


def i_a11y_voice(p):
    """Vorlesen — Text, der auch klingt."""
    for i, (x, y, h) in enumerate(((7, 28, 8), (17, 21, 22), (27, 15, 34),
                                   (37, 23, 18), (47, 30, 5))):
        p.solid(x, y, 6, h, GREEN if i == 2 else p.color)


# Zusammenarbeit
def i_team_handover(p):
    """Uebergabe — der Moment, in dem Wissen den Besitzer wechselt."""
    p.rect(4, 20, 18, 20)
    p.rect(38, 20, 18, 20)
    p.line(24, 30, 34, 30)
    p.poly([(32, 25), (38, 30), (32, 35)], close=True, fill=GREEN)


def i_team_review(p):
    """Review — jemand schaut wirklich hin."""
    p.rect(5, 19, 27, 25)
    p.poly([(11, 31), (17, 37), (27, 25)])
    p.circle(41, 25, 10)
    p.line(48, 32, 51, 35)
    p.dot(41, 25, 3.4, GREEN)


def i_team_fork(p):
    """Weggabelung — hier faellt die Entscheidung."""
    p.line(5, 31, 23, 31)
    p.poly([(23, 31), (37, 17), (51, 17)])
    p.poly([(23, 31), (37, 45), (51, 45)])
    p.dot(23, 31, 4.4, GREEN)


# Qualitaet
def i_debt_pile(p):
    """Technische Schuld — der Stapel, der schief steht."""
    p.rect(8, 38, 34, 10)
    p.rect(13, 27, 31, 10)
    p.rect(18, 16, 28, 10)
    p.rect(24, 5, 24, 10, fill=GREEN)


def i_refactor(p):
    """Aufraeumen — aus dem Knaeuel wird eine Linie."""
    p.poly([(7, 19), (23, 43), (7, 43), (23, 19)], close=True)
    p.line(31, 31, 50, 31)
    p.solid(45, 28, 6, 6, GREEN)


def i_legacy_crack(p):
    """Altsystem — die Schichten halten, aber es zieht sich ein Riss durch."""
    p.rect(6, 13, 44, 11)
    p.rect(6, 28, 44, 11)
    p.rect(6, 43, 40, 8)
    for x, y in ((26, 13), (23, 28), (28, 43)):
        p.solid(x, y, 5, 5, GREEN)


def i_test_net(p):
    """Testabdeckung — ehrlich: ein Feld ist wirklich geprueft."""
    p.rect(6, 11, 44, 36)
    p.line(21, 11, 21, 47)
    p.line(35, 11, 35, 47)
    p.line(6, 23, 50, 23)
    p.line(6, 35, 50, 35)
    p.solid(21, 23, 14, 12, GREEN)


# ── Dialekt B — Pixel-Art als String-Art ──────────────────────────────────
# Ein Motiv ist im Code direkt zu erkennen: '#' Signal-Gruen, 'o' dunkleres
# Gruen, '+' helles Gruen, 'x' Tinte, '.' leer. Raster 8x8 auf 6er-Zellen ab
# (4,4) — die rechte untere Ecke bleibt damit der Signatur.
_PX_PALETTE = {"#": "green", "o": "deep", "+": "soft", "x": "ink"}


def _pixels(p, art, cell=6.0, ox=4.0, oy=4.0):
    tone = {"green": GREEN, "deep": GREEN_DEEP, "soft": GREEN_SOFT, "ink": p.color}
    for r, row in enumerate(art):
        for c, ch in enumerate(row):
            key = _PX_PALETTE.get(ch)
            if key:
                p.solid(ox + c * cell, oy + r * cell, cell, cell, tone[key])


def i_pixel_dino(p):
    _pixels(p, [".....###",
                ".....#x#",
                ".....###",
                ".#...####"[:8],
                ".#######",
                ".######.",
                "..####..",
                "..#..#.."])


def i_pixel_ghost(p):
    _pixels(p, ["..####..",
                ".######.",
                "##x##x##",
                "########",
                "########",
                "########",
                "########",
                "#.#..#.#"])


def i_pixel_alien(p):
    _pixels(p, ["..#..#..",
                "...##...",
                "..####..",
                ".#x##x#.",
                "########",
                "#.####.#",
                "#.#..#.#",
                "..#..#.."])


def i_pixel_robot(p):
    _pixels(p, ["...##...",
                "..####..",
                ".######.",
                "#x####x#",
                ".######.",
                ".#.##.#.",
                ".######.",
                ".##..##."])


def i_pixel_cat(p):
    _pixels(p, ["#......#",
                "##....##",
                "########",
                "#x#..#x#",
                "########",
                "#.####.#",
                "########",
                ".#....#."])


def i_pixel_mushroom(p):
    _pixels(p, ["..####..",
                ".######.",
                "##+##+##",
                "########",
                ".oooooo.",
                "...oo...",
                "...oo...",
                "..oooo.."])


def i_pixel_bug(p):
    _pixels(p, ["#..##..#",
                ".#.##.#.",
                "..####..",
                ".#x##x#.",
                "########",
                ".######.",
                "#.####.#",
                ".#....#."])


def i_pixel_skull(p):
    _pixels(p, [".######.",
                "########",
                "#xx##xx#",
                "########",
                ".###x##.",
                ".######.",
                "..#x#x#.",
                "........"])


def i_pixel_crown(p):
    _pixels(p, [".#..#..#",
                ".#.##.#.",
                ".#.##.##",
                ".#######",
                ".#o#o#o#",
                ".#######",
                "........",
                "........"])


def i_pixel_coffee(p):
    _pixels(p, ["..o.o...",
                "...o.o..",
                "........",
                "#######.",
                "#+++++#o",
                "#+++++#o",
                "#######.",
                ".#####.."])


HEAVY = [
    ("grid", "Grid", i_grid), ("cursor", "Cursor", i_cursor), ("maze", "Labyrinth", i_maze),
    ("terminal", "Terminal", i_terminal), ("chat", "Chat", i_chat), ("spark", "Spark", i_spark),
    ("arrow", "Pfeil", i_arrow), ("user", "Person", i_user), ("gear", "Zahnrad", i_gear),
    ("folder", "Ordner", i_folder), ("lightbulb", "Glühbirne", i_lightbulb),
    ("code", "Code", i_code), ("calendar", "Kalender", i_calendar), ("clock", "Uhr", i_clock),
    ("chart", "Diagramm", i_chart), ("target", "Ziel", i_target), ("lock", "Schloss", i_lock),
    ("server", "Server", i_server), ("mobile", "Mobil", i_mobile), ("monitor", "Monitor", i_monitor),
    ("search", "Suche", i_search), ("mail", "Mail", i_mail), ("warning", "Warnung", i_warning),
    ("check", "Haken", i_check), ("close", "Kreuz", i_close), ("plus", "Plus", i_plus),
    ("external", "Extern", i_external), ("rocket", "Rakete", i_rocket), ("pen", "Stift", i_pen),
    ("trash", "Papierkorb", i_trash), ("filter", "Filter", i_filter), ("book", "Buch", i_book),
    ("bell", "Glocke", i_bell), ("phone", "Telefonhoerer", i_phone),
    ("callback", "Rueckruf", i_callback), ("form", "Formular", i_form),
]
PIXEL = [
    ("cucumber", "Pixel-Gurke", i_cucumber), ("pixel-heart", "Pixel-Herz", i_pixel_heart),
    ("pixel-star", "Pixel-Stern", i_pixel_star), ("pixel-disk", "Pixel-Diskette", i_pixel_disk),
    ("pixel-dino", "Pixel-Dino", i_pixel_dino), ("pixel-ghost", "Pixel-Geist", i_pixel_ghost),
    ("pixel-alien", "Pixel-Alien", i_pixel_alien), ("pixel-robot", "Pixel-Roboter", i_pixel_robot),
    ("pixel-cat", "Pixel-Katze", i_pixel_cat), ("pixel-mushroom", "Pixel-Pilz", i_pixel_mushroom),
    ("pixel-bug", "Pixel-Kaefer", i_pixel_bug), ("pixel-skull", "Pixel-Totenkopf", i_pixel_skull),
    ("pixel-crown", "Pixel-Krone", i_pixel_crown), ("pixel-coffee", "Pixel-Kaffee", i_pixel_coffee),
]

# ── Dialekt C — Konzept ───────────────────────────────────────────────────
# Abstrakter als die funktionalen Motive: Metaphern statt Gegenstaende, mehr
# Signal-Gruen, gern eine Prise Pixel. Gleiches Raster, gleiche Strichstaerke.

# Daten
def i_data_stream(p):
    """Datenstrom — verebbende Bloecke, gruene Spitze."""
    for y, cols in ((12, 5), (29, 4), (46, 3)):
        for i in range(cols):
            p.solid(8 + i * 8, y, 6, 6, p.color)
        p.solid(8 + cols * 8, y, 6, 6, GREEN)


def i_data_cluster(p):
    """Cluster — Streuung, die eine Huelle bekommt."""
    p.poly([(12, 18), (34, 8), (52, 24), (46, 46), (20, 48)], close=True)
    for cx, cy in ((22, 22), (32, 30), (24, 38), (38, 20), (40, 36)):
        p.dot(cx, cy, 3, p.color)
    p.dot(32, 30, 3.6, GREEN)


def i_data_pipe(p):
    """Pipeline — drei Stufen, eine traegt."""
    p.rect(5, 25, 15, 14)
    p.rect(25, 25, 15, 14, fill=GREEN)
    p.rect(45, 25, 15, 14)
    p.line(20, 32, 25, 32)
    p.line(40, 32, 45, 32)


def i_data_layers(p):
    """Schichten — vom Rohbestand zur Auswertung."""
    p.rect(6, 39, 45, 10)
    p.rect(11, 26, 35, 10)
    p.rect(16, 13, 25, 10, fill=GREEN)


def i_data_noise(p):
    """Rauschen und Signal — die gerade Linie ist die Aussage."""
    p.poly([(6, 28), (13, 14), (20, 32), (27, 12), (34, 30), (41, 16), (48, 32), (55, 18)])
    p.line(6, 46, 58, 46)
    p.solid(52, 43, 6, 6, GREEN)


# KI
def i_ai_grid(p):
    """Denkraster — ein Feld, in dem einzelne Zellen zuenden."""
    p.rect(8, 8, 43, 43)
    for cx, cy in ((15, 15), (30, 15), (22, 23), (37, 30), (15, 37), (30, 44)):
        p.solid(cx, cy, 6, 6, p.color)
    p.solid(37, 15, 6, 6, GREEN)
    p.solid(22, 37, 6, 6, GREEN)


def i_ai_neuron(p):
    """Neuron — ein Kern, der ausstrahlt."""
    for x, y in ((11, 15), (53, 15), (7, 32), (57, 32), (15, 51), (49, 51)):
        p.line(32, 32, x, y)
        p.dot(x, y, 3, p.color)
    p.circle(32, 32, 9, fill=GREEN)


def i_ai_prompt(p):
    """Prompt — die Frage, die den Funken ausloest."""
    p.poly([(8, 10), (56, 10), (56, 40), (26, 40), (16, 52), (16, 40), (8, 40)], close=True)
    p.poly([(32, 15), (35, 22), (43, 25), (35, 28), (32, 35), (29, 28), (21, 25), (29, 22)],
           close=True, fill=GREEN)


def i_ai_model(p):
    """Modell — viele Eingaben, eine Antwort. Bewusst duenn geknuepft,
    ein voll vernetztes Raster wird bei 48 px zum Fleck."""
    for a in (16, 48):
        for b in (12, 32, 52):
            p.line(10, a, 32, b)
    for b in (12, 32, 52):
        p.line(32, b, 54, 32)
    for y in (16, 48):
        p.dot(10, y, 3.4, p.color)
    for y in (12, 52):
        p.dot(32, y, 3.4, p.color)
    p.dot(32, 32, 4.2, GREEN)
    p.dot(54, 32, 3.8, p.color)


def i_ai_drift(p):
    """Drift — was fest begann und sich unterwegs aufloest."""
    p.solid(6, 24, 18, 18, p.color)
    for x, y, g in ((28, 26, 0), (34, 34, 1), (30, 44, 0), (40, 20, 0),
                    (44, 36, 1), (50, 28, 0), (55, 45, 0)):
        p.solid(x, y, 6, 6, GREEN if g else p.color)


# Workshop
def i_ws_sticky(p):
    """Haftnotizen — drei Gedanken, einer sticht heraus."""
    p.rect(4, 6, 23, 23)
    p.rect(16, 17, 23, 23, fill=GREEN)
    p.rect(28, 28, 23, 23)


def i_ws_vote(p):
    """Punktabfrage — wo die Klebepunkte landen."""
    p.line(6, 51, 50, 51)
    for i, n in enumerate((3, 5, 2)):
        x = 14 + i * 17
        for k in range(n):
            y = 44 - k * 9
            p.dot(x, y, 3.6, GREEN if i == 1 else p.color)


def i_ws_timebox(p):
    """Timebox — die Zeit laeuft, sichtbar."""
    p.poly([(13, 9), (51, 9), (32, 32)], close=True)
    p.poly([(32, 32), (51, 55), (13, 55)], close=True)
    p.solid(26, 46, 6, 6, GREEN)
    p.solid(32, 46, 6, 6, GREEN)


def i_ws_diverge(p):
    """Doppel-Diamant — erst oeffnen, dann entscheiden."""
    p.poly([(4, 32), (18, 16), (32, 32), (18, 48)], close=True)
    p.poly([(32, 32), (46, 16), (60, 32), (46, 48)], close=True, fill=GREEN)


def i_ws_agenda(p):
    """Agenda — Bloecke an einer Zeitachse."""
    p.line(14, 8, 14, 56)
    for i, y in enumerate((12, 26, 40)):
        p.rect(22, y, 32, 10, fill=GREEN if i == 1 else None)
        p.dot(14, y + 5, 3.4, p.color)


# Web
def i_web_layout(p):
    """Layout — Kopf, Spalte, Inhalt."""
    p.rect(6, 8, 45, 42)
    p.line(6, 20, 51, 20)
    p.solid(11, 26, 11, 20, GREEN)
    for y in (26, 34, 42):
        p.line(28, y + 2, 46, y + 2)


def i_web_responsive(p):
    """Responsiv — dasselbe, drei Breiten."""
    p.rect(5, 13, 28, 38)
    p.rect(37, 19, 14, 26)
    p.rect(54, 26, 6, 12, fill=GREEN)


def i_web_focus(p):
    """Fokus — der Ring, der niemals fehlen darf."""
    p.solid(21, 23, 22, 18, GREEN)
    p.rect(13, 15, 38, 34)


def i_web_speed(p):
    """Tempo — der Block ist schon da, die Spuren holen auf."""
    p.solid(42, 22, 18, 20, GREEN)
    p.line(4, 26, 38, 26)
    p.line(12, 32, 38, 32)
    p.line(20, 38, 38, 38)


def i_web_deploy(p):
    """Ausliefern — das Paket verlaesst den Boden."""
    p.rect(14, 32, 30, 18)
    p.line(29, 32, 29, 18)
    p.poly([(19, 24), (29, 12), (39, 24)], close=True, fill=GREEN)
    p.line(8, 56, 48, 56)


# Prototyping
def i_proto_loop(p):
    """Iteration — die Runde, die man mehrfach dreht."""
    p.poly([(44, 14), (14, 14), (14, 50), (50, 50), (50, 26)])
    p.poly([(38, 8), (48, 14), (38, 20)], close=True, fill=GREEN)


def i_proto_clickpath(p):
    """Klickpfad — der Weg, den jemand wirklich nimmt."""
    pts = [(10, 50), (23, 30), (38, 40), (50, 14)]
    p.poly(pts)
    for i, (x, y) in enumerate(pts):
        p.dot(x, y, 3.4, GREEN if i == 0 else p.color)
    p.poly([(50, 14), (50, 34), (56, 28), (60, 36)])


def i_proto_fidelity(p):
    """Detailgrad — von grob nach fein."""
    for x, y in ((6, 22), (14, 22), (6, 30), (14, 30)):
        p.solid(x, y, 6, 6, p.color)
    p.rect(26, 22, 14, 14)
    p.solid(46, 22, 14, 14, GREEN)


def i_proto_throwaway(p):
    """Wegwerf-Prototyp — gebaut, um verworfen zu werden.
    Flugbahn und Behaelter haengen zusammen, sonst zerfaellt das Motiv."""
    p.rect(4, 8, 15, 15)
    p.poly([(21, 13), (31, 8), (39, 15), (42, 27)])
    p.line(26, 32, 54, 32)
    p.poly([(30, 32), (33, 50), (47, 50), (50, 32)], close=True)
    p.solid(37, 39, 6, 6, GREEN)



CONCEPT = [
    ("data-stream", "Datenstrom", i_data_stream), ("data-cluster", "Cluster", i_data_cluster),
    ("data-pipe", "Pipeline", i_data_pipe), ("data-layers", "Schichten", i_data_layers),
    ("data-noise", "Rauschen", i_data_noise),
    ("ai-grid", "Denkraster", i_ai_grid), ("ai-neuron", "Neuron", i_ai_neuron),
    ("ai-prompt", "Prompt", i_ai_prompt), ("ai-model", "Modell", i_ai_model),
    ("ai-drift", "Drift", i_ai_drift),
    ("ws-sticky", "Haftnotizen", i_ws_sticky), ("ws-vote", "Punktabfrage", i_ws_vote),
    ("ws-timebox", "Timebox", i_ws_timebox), ("ws-diverge", "Doppel-Diamant", i_ws_diverge),
    ("ws-agenda", "Agenda", i_ws_agenda),
    ("web-layout", "Layout", i_web_layout), ("web-responsive", "Responsiv", i_web_responsive),
    ("web-focus", "Fokus", i_web_focus), ("web-speed", "Tempo", i_web_speed),
    ("web-deploy", "Deploy", i_web_deploy),
    ("proto-loop", "Iteration", i_proto_loop), ("proto-clickpath", "Klickpfad", i_proto_clickpath),
    ("proto-fidelity", "Detailgrad", i_proto_fidelity),
    ("proto-throwaway", "Wegwerf-Prototyp", i_proto_throwaway),
    ("ops-pulse", "Monitoring", i_ops_pulse), ("ops-incident", "Stoerung", i_ops_incident),
    ("ops-rollback", "Rollback", i_ops_rollback), ("ops-scale", "Skalieren", i_ops_scale),
    ("ops-backup", "Sicherung", i_ops_backup),
    ("sec-encrypt", "Verschluesseln", i_sec_encrypt),
    ("sec-surface", "Angriffsflaeche", i_sec_surface), ("sec-key", "Berechtigung", i_sec_key),
    ("a11y-contrast", "Kontrast", i_a11y_contrast),
    ("a11y-keyboard", "Tastaturpfad", i_a11y_keyboard), ("a11y-voice", "Vorlesen", i_a11y_voice),
    ("team-handover", "Uebergabe", i_team_handover), ("team-review", "Review", i_team_review),
    ("team-fork", "Weggabelung", i_team_fork),
    ("debt-pile", "Technische Schuld", i_debt_pile), ("refactor", "Aufraeumen", i_refactor),
    ("legacy-crack", "Altsystem", i_legacy_crack), ("test-net", "Testabdeckung", i_test_net),
]

ALL = HEAVY + PIXEL + CONCEPT
BY_SLUG = {s: (lbl, fn) for s, lbl, fn in ALL}
DIALECT = {**{s: "heavy" for s, _, _ in HEAVY},
           **{s: "pixel" for s, _, _ in PIXEL},
           **{s: "concept" for s, _, _ in CONCEPT}}


# Die Signatur: derselbe 6x6-Punkt unten rechts, den auch die 462 Katalog-
# Icons tragen. Sie wird zentral beim Zeichnen gesetzt — nicht in den
# einzelnen Motiven —, damit kein kuenftiges Icon sie vergessen kann.
# Ganzzahlig, damit das Kernset denselben Quelltext ausgibt wie der Katalog
# (rect x="54" …, nicht x="54.0" …) — eine Marke, eine Schreibweise.
SIGNATURE = (54, 54, 6, 6)
RESERVED = (52.0, 52.0)          # Faustregel fuer Entwuerfe: Ecke frei lassen


def _sign(pen):
    pen.solid(*SIGNATURE, GREEN)


def draw(slide, slug, x, y, size, color):
    pen = Pen(slide, x, y, size, color)
    BY_SLUG[slug][1](pen)
    _sign(pen)


def to_svg(slug, color="currentColor"):
    lbl, fn = BY_SLUG[slug]
    pen = SvgPen(color)
    fn(pen)
    _sign(pen)
    return pen.svg(lbl)


class BoundsPen:
    """Zeichnet nichts, sammelt nur Geometrie — fuer die Pruefung, ob ein Motiv
    in die fuer die Signatur reservierte Ecke ragt. Bewusst exakt statt per
    Huellrechteck: ein Stern oder ein Kreis beruehrt seine Bounding-Box-Ecke nie."""

    HALF = 2.0                      # halbe Strichstaerke

    def __init__(self):
        self.color = "ink"
        self.segs = []              # Strecken (Umrisse)
        self.discs = []             # (cx, cy, r, gefuellt)
        self.boxes = []             # gefuellte Rechtecke
        self.polys = []             # gefuellte Vielecke

    def line(self, x1, y1, x2, y2):
        self.segs.append((x1, y1, x2, y2))

    def _rect_segs(self, x, y, w, h):
        self.segs += [(x, y, x + w, y), (x + w, y, x + w, y + h),
                      (x + w, y + h, x, y + h), (x, y + h, x, y)]

    def rect(self, x, y, w, h, fill=None):
        self._rect_segs(x, y, w, h)
        if fill is not None:
            self.boxes.append((x, y, x + w, y + h))

    def solid(self, x, y, w, h, fill):
        self.boxes.append((x, y, x + w, y + h))

    def circle(self, cx, cy, r, fill=None):
        self.discs.append((cx, cy, r, fill is not None))

    def dot(self, cx, cy, r, fill):
        self.discs.append((cx, cy, r, True))

    def poly(self, pts, close=False, fill=None):
        pairs = list(zip(pts, pts[1:])) + ([(pts[-1], pts[0])] if close else [])
        self.segs += [(a[0], a[1], b[0], b[1]) for a, b in pairs]
        if fill is not None:
            self.polys.append(list(pts))


def _seg_hits_rect(x1, y1, x2, y2, r):
    """Liang-Barsky: schneidet die Strecke das Rechteck?"""
    rx0, ry0, rx1, ry1 = r
    dx, dy = x2 - x1, y2 - y1
    p = (-dx, dx, -dy, dy)
    q = (x1 - rx0, rx1 - x1, y1 - ry0, ry1 - y1)
    t0, t1 = 0.0, 1.0
    for pi, qi in zip(p, q):
        if pi == 0:
            if qi < 0:
                return False
        else:
            t = qi / pi
            if pi < 0:
                if t > t1:
                    return False
                t0 = max(t0, t)
            else:
                if t < t0:
                    return False
                t1 = min(t1, t)
    return True


def _disc_hits_rect(cx, cy, rad, filled, r):
    from math import hypot
    rx0, ry0, rx1, ry1 = r
    nx, ny = max(rx0, min(cx, rx1)), max(ry0, min(cy, ry1))
    near = hypot(cx - nx, cy - ny)
    if filled:
        return near <= rad
    far = max(hypot(cx - x, cy - y) for x in (rx0, rx1) for y in (ry0, ry1))
    return near <= rad + BoundsPen.HALF and far >= rad - BoundsPen.HALF


def _point_in_poly(px, py, pts):
    inside = False
    j = len(pts) - 1
    for i in range(len(pts)):
        xi, yi = pts[i]; xj, yj = pts[j]
        if (yi > py) != (yj > py) and px < (xj - xi) * (py - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def clearance_violations():
    """Motive, die in die reservierte Ecke ragen — leere Liste heisst sauber."""
    # Gesperrt ist genau die Flaeche des Punktes (54..60), geweitet um die halbe
    # Strichstaerke der Motive — nicht der ganze Quadrant: der Katalog reicht
    # selbst bis y=54, nur eben nicht in die Ecke.
    h = BoundsPen.HALF
    sx, sy, sw, sh = SIGNATURE
    zone = (sx - h, sy - h, sx + sw + h, sy + sh + h)
    bad = []
    for slug, label, fn in ALL:
        pen = BoundsPen()
        fn(pen)
        hit = (any(_seg_hits_rect(*seg, zone) for seg in pen.segs)
               or any(_disc_hits_rect(*d, zone) for d in pen.discs)
               or any(b[2] > zone[0] and b[3] > zone[1] and b[0] < zone[2] and b[1] < zone[3]
                      for b in pen.boxes)
               or any(_point_in_poly(cx, cy, poly)
                      for poly in pen.polys
                      for cx in (zone[0], zone[2]) for cy in (zone[1], zone[3])))
        if hit:
            bad.append((slug, label))
    return bad


def wordmark(slide, x, y, w, color, green=GREEN):
    """Wordmark as live text: lowercase 'nozilla' plus the green period."""
    from pptx.util import Inches
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w * 1.25), Inches(w * 0.34))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    for txt, col in (("nozilla", color), (".", green)):
        r = para.add_run()
        r.text = txt
        r.font.name = "Zilla Slab"
        r.font.bold = True
        r.font.size = Pt(w * 34)
        r.font.color.rgb = col
    return tb
